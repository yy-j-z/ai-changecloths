from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt


SOURCE = Path(r"C:\Users\jyy\OneDrive\Desktop\需求分析文档模板_形镜AI变装.docx")
OUT = Path(r"C:\课程\软件工程-大三上\ai换装\deliverables\需求分析文档_形镜AI变装_含图1图2.docx")
FIG1 = Path(r"C:\课程\软件工程-大三上\ai换装\deliverables\图1_系统功能架构图.png")
FIG2 = Path(r"C:\课程\软件工程-大三上\ai换装\deliverables\图2_系统总体技术架构图.png")


def paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def set_run_font(run, name: str, size: float, bold: bool = False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def replace_placeholder(paragraph: Paragraph, image_path: Path, caption: str, page_break_before: bool = False):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.78))
    drawing = run._element.find(".//w:drawing", run._element.nsmap)
    if drawing is not None:
        doc_pr = drawing.find(".//wp:docPr", drawing.nsmap)
        if doc_pr is not None:
            doc_pr.set("descr", caption)

    cap = paragraph_after(paragraph)
    cap.style = "Caption" if "Caption" in [s.name for s in paragraph._parent.part.document.styles] else paragraph.style
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = False
    cap_run = cap.add_run(caption)
    set_run_font(cap_run, "宋体", 10.5)


doc = Document(SOURCE)

figure1_para = next(p for p in doc.paragraphs if p.text.startswith("📊 图 1 系统功能架构图"))
figure2_para = next(p for p in doc.paragraphs if p.text.startswith("📊 图 2 系统总体技术架构图"))
heading_para = next(p for p in doc.paragraphs if p.text.strip() == "1.8 功能架构图")
heading_para.paragraph_format.keep_with_next = True

replace_placeholder(figure1_para, FIG1, "图 1  系统功能架构图")
replace_placeholder(figure2_para, FIG2, "图 2  系统总体技术架构图", page_break_before=True)

settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
